import openpyxl, hashlib

files = {"c9d9": "c9d9.xlsx", "642a": "642a.xlsx", "dc53": "dc53.xlsx"}
data = {}

for k, p in files.items():
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    print(f"{k}: シート={wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    data[k] = rows
    digest = hashlib.sha256(repr(rows).encode()).hexdigest()[:12]
    print(f"      {len(rows)}行 x {len(rows[0]) if rows else 0}列  値ハッシュ={digest}")

print()
for k, rows in data.items():
    print(f"--- {k} 先頭3行 ---")
    for r in rows[:3]:
        print("   ", r)

def diff(a, b):
    ra, rb = data[a], data[b]
    n = 0
    for i in range(min(len(ra), len(rb))):
        if ra[i] != rb[i]:
            if n < 5:
                print(f"  行{i+1}:")
                print(f"    {a}: {ra[i]}")
                print(f"    {b}: {rb[i]}")
            n += 1
    print(f"{a} vs {b}: 差分 {n} 行 (行数 {len(ra)} / {len(rb)})\n")

print()
diff("c9d9", "642a")
diff("dc53", "642a")
diff("c9d9", "dc53")