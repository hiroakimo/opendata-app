"""目黒区 BODIK 月次更新の確認と投入SQLの抜き出し（D1には書かない）

対象はマニフェスト（sync.py が書く bodik.jsonl）から自動で解決する。
  --month 2026-10         その月の 1歳階級・5歳階級
  --extra 2024-12-01      追加で取り込む基準日（差し替え・欠測の補完など。複数可）
  --hold  2017-07-01      取り込まない基準日（既定: 2017-07-01。複数可）

手順（リポジトリ直下で）:
  1. python tools/meguro_monthly.py check   --month 2026-10 [--extra ...]
       取得済みファイルの中身を確認する（読み取りのみ）
  2. python pipeline/wards/meguro/bodik/normalize.py --manifest ... （SQL生成）
  3. python tools/meguro_monthly.py extract --month 2026-10 [--extra ...]
       生成SQLから対象の基準日だけを抜き出し、投入と確認のコマンドを表示する
検証に1つでも失敗したら、extract は何も書き出さない。
"""
import argparse
import csv
import io
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path("work/meguro/bodik")
MANIFEST = BASE / "manifest/bodik.jsonl"
SRC = BASE / "src"
CHECK = BASE / "check"
SQL = BASE / "sql"
DEFAULT_HOLD = ["2017-07-01"]
DATASET_KEY = {"5y": "meguro_5y", "1y": "meguro_1y"}
OBS_TABLE = {"5y": "observations_5y", "1y": "observations_1y"}
OBS_FILE = {"5y": "10_5y_{y}.sql", "1y": "11_1y_{y}.sql"}

FILE_RE = re.compile(r"^131105_population_(?:(1y|5y)_)?(\d{4})(\d{2})(\d{2})\.(csv|xlsx)$")
HEX64 = re.compile(r"^[0-9a-f]{64}$")
DATE_RE = re.compile(r"(令和|平成)\S*年|\d{4}年\s*\d{1,2}月|\d{4}[-/]\d{1,2}[-/]\d{1,2}")
KANSUJI = str.maketrans("一二三四五六七八九", "123456789")

out_lines = []
errors = []


def out(s=""):
    print(s)
    out_lines.append(s)


def ng(msg):
    errors.append(msg)
    out(f"  NG  {msg}")


# ---------------------------------------------------------------- 共通

def month_date(s):
    m = re.fullmatch(r"(\d{4})-(\d{2})(?:-01)?", s)
    if not m:
        raise SystemExit(f"--month は YYYY-MM で指定してください: {s}")
    return f"{m.group(1)}-{m.group(2)}-01"


def day_date(s):
    if not re.fullmatch(r"\d{4}-\d{2}-01", s):
        raise SystemExit(f"基準日は YYYY-MM-01 で指定してください: {s}")
    return s


def shift_month(d, k):
    y, m = int(d[:4]), int(d[5:7])
    t = y * 12 + (m - 1) + k
    return f"{t // 12:04d}-{t % 12 + 1:02d}-01"


def sha256_of(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _strings(v, key=""):
    if isinstance(v, str):
        yield key, v
    elif isinstance(v, dict):
        for k, x in v.items():
            yield from _strings(x, k)
    elif isinstance(v, list):
        for x in v:
            yield from _strings(x, key)


def load_manifest():
    """ファイル名ごとに、マニフェストの最後の記録（= 現行版）を返す。"""
    if not MANIFEST.exists():
        raise SystemExit(f"マニフェストがありません: {MANIFEST}")
    latest = {}
    with open(MANIFEST, encoding="utf-8") as f:
        for n, line in enumerate(f, 1):
            if not line.strip():
                continue
            rec = json.loads(line)
            strs = list(_strings(rec))
            fname = None
            for _, v in strs:
                name = re.split(r"[\\/]", v)[-1]
                if FILE_RE.match(name):
                    fname = name
                    break
            if not fname:
                continue
            sha = rec.get("sha256") if isinstance(rec.get("sha256"), str) else None
            if not (sha and HEX64.match(sha)):
                cands = {v for k, v in strs if HEX64.match(v) and "supersed" not in k}
                sha = cands.pop() if len(cands) == 1 else None
            if not sha:
                continue  # MISSING などの記録
            sup = rec.get("supersedes")
            m = FILE_RE.match(fname)
            latest[fname] = {
                "file": fname,
                "sha": sha,
                "supersedes": sup if isinstance(sup, str) and HEX64.match(sup) else None,
                "status": rec.get("status") or rec.get("change") or "",
                "grain": m.group(1) or "5y",
                "date": f"{m.group(2)}-{m.group(3)}-{m.group(4)}",
                "ext": m.group(5),
                "line": n,
            }
    return latest


def resolve_targets(args):
    month = month_date(args.month)
    extras = [day_date(d) for d in args.extra]
    hold = set(day_date(d) for d in (args.hold or DEFAULT_HOLD))
    for d in [month] + extras:
        if d in hold:
            raise SystemExit(f"{d} は --hold に含まれています")
    latest = load_manifest()
    by_date = defaultdict(list)
    for e in latest.values():
        by_date[e["date"]].append(e)

    targets = []
    out(f"■ 対象（マニフェスト: {MANIFEST}）")
    for d in [month] + extras:
        found = sorted(by_date.get(d, []), key=lambda e: e["grain"])
        if d == month:
            grains = {e["grain"] for e in found}
            for g in ("1y", "5y"):
                if g not in grains:
                    ng(f"{d} の {g} がマニフェストにありません（sync.py の本実行は済んでいますか）")
        elif not found:
            ng(f"{d} がマニフェストにありません")
        for e in found:
            e["role"] = "month" if d == month else "extra"
            sup = f" / 前の版 {e['supersedes'][:12]}" if e["supersedes"] else ""
            out(f"  {e['role']:5s} {e['grain']} {d}  {e['file']}  sha {e['sha'][:12]}{sup}")
            targets.append(e)
    out(f"  保留: {', '.join(sorted(hold))}")
    out()
    return month, extras, hold, targets


# ---------------------------------------------------------------- check

def read_csv_rows(p):
    raw = p.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        text, enc = raw.decode("utf-16"), "utf-16"
    else:
        for enc in ("utf-8-sig", "cp932"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("文字コードを判定できません")
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    return rows, enc


def read_xlsx_rows(p):
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        sheets.append((ws.title, rows))
    wb.close()
    return sheets


def cell_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str) and DATE_RE.search(v):
        return v.strip()[:40]
    return None


def num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    return int(str(v).replace(",", "").strip())


def load_table(e, path):
    """(見出し, データ行) を返す。構造の要約も出力する。"""
    if e["ext"] == "csv":
        rows, enc = read_csv_rows(path)
        widths = Counter(len(r) for r in rows)
        out(f"  文字コード {enc} / 行 {len(rows)} / 列数の分布 {dict(widths)}")
        if len(widths) != 1:
            ng(f"{e['file']}: 行によって列数が違います")
        return rows[0], rows[1:]
    sheets = read_xlsx_rows(path)
    out(f"  シート {[s for s, _ in sheets]}")
    for title, rows in sheets:
        width = max((max(i for i, v in enumerate(r) if v is not None) + 1 for r in rows), default=0)
        out(f"  [{title}] 行 {len(rows)} / 列 {width}")
    if len(sheets) != 1:
        ng(f"{e['file']}: シートが {len(sheets)} 枚あります（これまでは1枚）")
    rows = sheets[0][1]
    return rows[0], rows[1:]


def check_dates(e, data):
    dates = Counter()
    for r in data:
        for v in r:
            d = cell_date(v)
            if d:
                dates[d] += 1
    out(f"  日付らしい値: {dict(dates.most_common(4))}")
    if set(dates) != {e["date"]}:
        ng(f"{e['file']}: 調査日が基準日 {e['date']} と一致しません")


def check_5y_internal(e, hdr, data):
    idx = {str(h): i for i, h in enumerate(hdr)}
    for need in ("地域名", "総人口", "男性", "女性"):
        if need not in idx:
            ng(f"{e['file']}: 見出し「{need}」がありません")
            return None
    m_cols = [i for i, h in enumerate(hdr) if str(h).endswith("の男性")]
    f_cols = [i for i, h in enumerate(hdr) if str(h).endswith("の女性")]
    others = [h for h in hdr[idx["女性"] + 1:] if not str(h).endswith(("の男性", "の女性"))]
    out(f"  町丁目 {len(data)} / 年齢別列 男{len(m_cols)} 女{len(f_cols)} / その他 {others}")
    bad, tot, table = 0, [0, 0, 0], {}
    for r in data:
        t, m, f = num(r[idx["総人口"]]), num(r[idx["男性"]]), num(r[idx["女性"]])
        sm, sf = sum(num(r[i]) for i in m_cols), sum(num(r[i]) for i in f_cols)
        tot[0] += t; tot[1] += m; tot[2] += f
        table[str(r[idx["地域名"]]).translate(KANSUJI)] = (t, m, f)
        if t != m + f or sm != m or sf != f:
            bad += 1
            if bad <= 5:
                out(f"    {r[idx['地域名']]}: 総{t} 男{m}(年齢計{sm}) 女{f}(年齢計{sf})")
    out(f"  内部整合の不一致 {bad} / 区合計 総人口 {tot[0]:,} 男 {tot[1]:,} 女 {tot[2]:,}")
    if bad:
        ng(f"{e['file']}: 内部整合の不一致 {bad} 町丁目")
    return table


def agg_1y(e, hdr, data):
    idx = {str(h): i for i, h in enumerate(hdr)}
    for need in ("町丁目", "総数", "男性", "女性"):
        if need not in idx:
            ng(f"{e['file']}: 見出し「{need}」がありません")
            return None
    agg = defaultdict(lambda: [0, 0, 0])
    for r in data:
        a = agg[str(r[idx["町丁目"]]).translate(KANSUJI)]
        a[0] += num(r[idx["総数"]]); a[1] += num(r[idx["男性"]]); a[2] += num(r[idx["女性"]])
    out(f"  町丁目 {len(agg)} / 行 {len(data)} / 区合計 総数 {sum(v[0] for v in agg.values()):,}")
    return {k: tuple(v) for k, v in agg.items()}


def compare_versions(e, old_path, new_hdr, new_data):
    old_rows, _ = read_csv_rows(old_path)
    old_hdr, old_data = old_rows[0], old_rows[1:]
    widths = Counter(len(r) for r in old_rows)
    out(f"  旧版 {e['supersedes'][:12]}: 行 {len(old_rows)} / 列数の分布 {dict(widths)}")
    out(f"  見出し: {'同一' if old_hdr == new_hdr else '異なる'}")
    so, sn = set(map(tuple, old_data)), set(map(tuple, new_data))
    out(f"  行の集合: 共通 {len(so & sn)} / 旧のみ {len(so - sn)} / 新のみ {len(sn - so)}")
    shown = 0
    for i, (ro, rn) in enumerate(zip(old_data, new_data), start=2):
        for j, (a, b) in enumerate(zip(ro, rn)):
            if a != b and shown < 10:
                name = new_hdr[j] if j < len(new_hdr) else f"列{j + 1}"
                out(f"    {i}行目 {rn[:2]} {name}: {a!r} → {b!r}")
                shown += 1


def cmd_check(args):
    month, extras, hold, targets = resolve_targets(args)
    pending = []
    tables = {}
    for e in targets:
        path = SRC / e["file"]
        out(f"■ {e['grain']} {e['date']}  {path}")
        if not path.exists():
            ng(f"手元にファイルがありません（sync.py に --keep-local {SRC} を付けて実行してください）")
            continue
        if sha256_of(path) != e["sha"]:
            ng(f"{e['file']}: 手元のファイルの sha256 がマニフェストと一致しません")
            continue
        hdr, data = load_table(e, path)
        check_dates(e, data)
        if e["grain"] == "5y":
            tables[(e["grain"], e["date"])] = check_5y_internal(e, hdr, data)
        else:
            tables[(e["grain"], e["date"])] = agg_1y(e, hdr, data)
        if e["supersedes"] and e["ext"] == "csv":
            old = CHECK / f"{e['supersedes']}.csv"
            if old.exists():
                compare_versions(e, old, hdr, data)
            else:
                pending.append(
                    f"npx wrangler r2 object get opendata-lake/raw/{e['supersedes']} --remote --file {old}")
        out()

    t1, t5 = tables.get(("1y", month)), tables.get(("5y", month))
    if t1 is not None and t5 is not None:
        out(f"■ {month} 1歳×5歳 突き合わせ（町丁目ごとの 総数/男性/女性）")
        diff = [k for k in set(t1) | set(t5) if t1.get(k) != t5.get(k)]
        for k in sorted(diff)[:10]:
            out(f"    {k}: 1歳 {t1.get(k)} / 5歳 {t5.get(k)}")
        out(f"  町丁目 1歳 {len(t1)} / 5歳 {len(t5)} / 不一致 {len(diff)}")
        if diff:
            ng(f"{month}: 1歳階級と5歳階級が {len(diff)} 町丁目で一致しません")
        out()

    dates5 = sorted({e["date"] for e in targets if e["grain"] == "5y"})
    around = sorted({shift_month(d, k) for d in dates5 for k in (-2, -1, 1, 2)} - set(dates5) - hold)
    out("■ D1 との連続性（次のクエリで、前後の月の区合計と比べてください）")
    for d in dates5:
        t = tables.get(("5y", d))
        if t:
            out(f"  {d}: 総人口 {sum(v[0] for v in t.values()):,}")
    out('  npx wrangler d1 execute tokyo-population --remote --command "SELECT reference_date, SUM(value) AS pop '
        "FROM observations_5y WHERE measure='population' AND age_class='total' AND sex='total' AND reference_date IN ("
        + ",".join(f"'{d}'" for d in around) + ') GROUP BY reference_date ORDER BY reference_date"')
    out()

    CHECK.mkdir(parents=True, exist_ok=True)
    report = CHECK / f"report_{month[:7].replace('-', '')}.txt"
    if pending:
        out("■ 旧版が手元にありません。次を実行してから、もう一度 check を実行してください。")
        for c in pending:
            out(f"  {c}")
    out(f"結果: NG {len(errors)} 件" + (" / 旧版の取得待ちあり" if pending else ""))
    report.write_text("\n".join(out_lines), encoding="utf-8")
    print(f"記録しました: {report}")
    return 1 if errors else (2 if pending else 0)


# ---------------------------------------------------------------- extract

def split_statements(path):
    buf = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\r\n")
            if not buf and (not line.strip() or line.lstrip().startswith("--")):
                continue
            buf.append(line)
            if line.rstrip().endswith(";"):
                yield buf
                buf = []
    if buf:
        ng(f"{path.name}: 終端の ';' がない文があります（{buf[0][:80]}）")


def parse_tuple(line):
    s = line.strip().rstrip(",;").strip()
    if not (s.startswith("(") and s.endswith(")")):
        raise ValueError(f"行の形式が想定外: {line[:80]}")
    s, vals, i, n = s[1:-1], [], 0, len(s) - 2
    while i < n:
        while i < n and s[i] == " ":
            i += 1
        if s[i] == "'":
            j, buf = i + 1, []
            while True:
                if s[j] == "'":
                    if j + 1 < n and s[j + 1] == "'":
                        buf.append("'"); j += 2
                        continue
                    break
                buf.append(s[j]); j += 1
            vals.append("".join(buf))
            i = j + 1
        else:
            j = s.find(",", i)
            j = n if j < 0 else j
            tok = s[i:j].strip()
            vals.append(None if tok == "NULL" else (float(tok) if "." in tok else int(tok)))
            i = j
        while i < n and s[i] == " ":
            i += 1
        if i < n:
            if s[i] != ",":
                raise ValueError(f"区切りが想定外: {line[:80]}")
            i += 1
    return vals


HEAD_RE = re.compile(r"^INSERT OR REPLACE INTO (\w+) \((.*)\) VALUES\s*$")
OBS_ROW_RE = re.compile(r"^\('[^']*', '[^']*', '(?P<date>\d{4}-\d{2}-\d{2})', .*'(?P<sha>[0-9a-f]{64})'\)[,;]\s*$")
DEL_RE = re.compile(r"^DELETE FROM (\w+) WHERE reference_date = '(\d{4}-\d{2}-\d{2})';\s*$")


def load_meta(hold):
    """20_meta.sql を読み、source_files の行（sha→(列dict, 行テキスト)）と、それ以外の文を返す。"""
    sf_head, sf_rows, others = None, {}, []
    hold_shas = set()
    for st in split_statements(SQL / "20_meta.sql"):
        m = HEAD_RE.match(st[0])
        if not m:
            ng(f"20_meta.sql に想定外の文があります: {st[0][:100]}")
            continue
        table, cols = m.group(1), [c.strip() for c in m.group(2).split(",")]
        if table == "source_files":
            sf_head = (st[0], cols)
            for line in st[1:]:
                rec = dict(zip(cols, parse_tuple(line)))
                sf_rows[rec["sha256"]] = (rec, line.rstrip().rstrip(",;"))
                if rec.get("reference_date") in hold:
                    hold_shas.add(rec["sha256"])
        else:
            others.append((table, st))
    if sf_head is None:
        ng("20_meta.sql に source_files がありません")
    return sf_head, sf_rows, others, hold_shas


def cmd_extract(args):
    month, extras, hold, targets = resolve_targets(args)
    out_dir = BASE / f"load_{month[:7].replace('-', '')}"
    sf_head, sf_rows, others, hold_shas = load_meta(hold)
    if errors:
        return finish_fail()
    cols = sf_head[1]

    # source_files の行を確認し、期待する行数を決める
    out("■ source_files")
    expect = {}
    old_shas = set()
    for e in targets:
        row = sf_rows.get(e["sha"])
        if not row:
            ng(f"{e['file']} {e['sha'][:12]} の行が 20_meta.sql にありません")
            continue
        rec = row[0]
        cur = rec.get("is_current", "?")
        out(f"  {e['grain']} {e['date']}  {e['sha'][:12]}  status={rec.get('status')} "
            f"is_current={cur} row_count={rec.get('row_count')}")
        if rec.get("status") != "ingested" or not rec.get("row_count"):
            ng(f"{e['file']}: 取り込み対象になっていません（status={rec.get('status')} "
               f"skip_reason={rec.get('skip_reason')}）")
        if "is_current" in rec and rec["is_current"] != 1:
            ng(f"{e['file']}: is_current が 1 ではありません")
        expect[(e["grain"], e["date"])] = (e["sha"], rec.get("row_count"))
        for s in (e["supersedes"], rec.get("supersedes")):
            if s and HEX64.match(s):
                old_shas.add(s)
    for s in sorted(old_shas):
        if s in sf_rows:
            r = sf_rows[s][0]
            out(f"  前の版 {s[:12]}  status={r.get('status')} is_current={r.get('is_current', '?')}"
                "  → 既存の行は保持し、is_current だけ 0 にする")
        else:
            out(f"  前の版 {s[:12]}  20_meta.sql に行なし → is_current の更新だけ行う")
    out()

    # 観測値
    outputs = {}
    for grain in ("5y", "1y"):
        want = {d: v for (g, d), v in expect.items() if g == grain}
        if not want:
            continue
        table = OBS_TABLE[grain]
        files = sorted({OBS_FILE[grain].format(y=d[:4]) for d in want})
        name = f"30_obs_{grain}.sql" if grain == "5y" else f"31_obs_{grain}.sql"
        out(f"■ {name}  ← {', '.join(files)}")
        kept, dels, rows = [], Counter(), Counter()
        for fn in files:
            if not (SQL / fn).exists():
                ng(f"{fn} がありません")
                continue
            for st in split_statements(SQL / fn):
                dm = DEL_RE.match(st[0])
                if dm:
                    if dm.group(1) != table:
                        ng(f"{fn}: 想定外のテーブルへの DELETE: {st[0]}")
                    elif dm.group(2) in want:
                        dels[dm.group(2)] += 1
                        kept.append(st)
                    continue
                if not st[0].startswith(f"INSERT OR REPLACE INTO {table} ("):
                    ng(f"{fn}: 想定外の文: {st[0][:100]}")
                    continue
                ms = [OBS_ROW_RE.match(r) for r in st[1:]]
                if not all(ms):
                    ng(f"{fn}: 行の形式が想定外の INSERT があります")
                    continue
                ds = {m.group("date") for m in ms}
                if len(ds) != 1:
                    ng(f"{fn}: 1つの INSERT に複数の基準日 {sorted(ds)}")
                    continue
                d = ds.pop()
                if d not in want:
                    continue
                if any(m.group("sha") != want[d][0] for m in ms):
                    ng(f"{d}: 想定外の source_sha256 を含む行があります")
                    continue
                rows[d] += len(ms)
                kept.append(st)
        for d, (sha, n) in sorted(want.items()):
            ok = dels[d] == 1 and rows[d] == n
            out(f"  {'OK' if ok else 'NG'}  {d}  DELETE {dels[d]} / 行 {rows[d]:,}（期待 {n:,}）/ sha {sha[:12]}")
            if not ok:
                errors.append(f"{name} {d}")
        outputs[name] = "\n".join("\n".join(st) for st in kept) + "\n"

    # メタデータ
    out("■ 32_meta.sql  ← 20_meta.sql")
    head_line = sf_head[0]
    new_rows = [sf_rows[e["sha"]][1] for e in targets if e["sha"] in sf_rows]
    parts = [head_line + "\n" + ",\n".join(new_rows) + ";"]
    keep_old = [sf_rows[s][1] for s in sorted(old_shas) if s in sf_rows]
    if keep_old:
        parts.append(head_line.replace("INSERT OR REPLACE", "INSERT OR IGNORE", 1)
                     + "\n" + ",\n".join(keep_old) + ";")
    if old_shas and "is_current" in cols:
        parts.append("UPDATE source_files SET is_current = 0 WHERE sha256 IN ("
                     + ", ".join(f"'{s}'" for s in sorted(old_shas)) + ");")
    out(f"  source_files: 対象 {len(new_rows)} 行 / 前の版 {len(old_shas)} 件")
    for table, st in others:
        text = "\n".join(st)
        if any(h in text for h in hold) or any(s in text for s in hold_shas):
            ng(f"{table}: 保留中の基準日を含むため、そのままは流せません")
            continue
        out(f"  {table}: {len(st) - 1} 行（そのまま採用）")
        parts.append(text)
    outputs["32_meta.sql"] = "\n".join(parts) + "\n"

    # 欠測の記録
    if extras:
        gap = ["-- 補完した基準日の、取り込み失敗の記録を外す"]
        for e in targets:
            if e["role"] == "extra":
                gap.append(f"DELETE FROM dataset_gaps WHERE dataset_key = '{DATASET_KEY[e['grain']]}' "
                           f"AND reference_date = '{e['date']}' AND kind = 'ingest_failed';")
        outputs["33_gap.sql"] = "\n".join(gap) + "\n"

    # 最終確認
    for name, text in outputs.items():
        if any(h in text for h in hold) or any(s in text for s in hold_shas):
            ng(f"{name}: 保留中の基準日が含まれています")
    if errors:
        return finish_fail()

    out_dir.mkdir(parents=True, exist_ok=True)
    out()
    for name, text in outputs.items():
        (out_dir / name).write_text(text, encoding="utf-8")
        out(f"書き出し: {out_dir / name}  {len(text.encode('utf-8')):,} bytes")
    print_commands(out_dir, outputs, expect)
    (out_dir / "commands.txt").write_text("\n".join(out_lines), encoding="utf-8")
    print(f"\n記録しました: {out_dir / 'commands.txt'}")
    return 0


def finish_fail():
    out(f"\n検証に失敗しました（{len(errors)} 件）。何も書き出していません。")
    return 1


def print_commands(out_dir, outputs, expect):
    q = 'npx wrangler d1 execute tokyo-population --remote --command "{}"'
    dates5 = sorted(d for g, d in expect if g == "5y")
    dates1 = sorted(d for g, d in expect if g == "1y")
    inlist = lambda ds: ",".join(f"'{d}'" for d in ds)

    out("\n■ 1. 投入（直前に巻き戻し地点を控える）")
    out("  npx wrangler d1 time-travel info tokyo-population")
    for name in outputs:
        out(f"  npx wrangler d1 execute tokyo-population --remote --file {out_dir / name}")

    out("\n■ 2. 投入後の確認")
    sel = ["SELECT '5y' AS g, reference_date, source_sha256, COUNT(*) AS n FROM observations_5y "
           f"WHERE reference_date IN ({inlist(dates5)}) GROUP BY reference_date, source_sha256"] if dates5 else []
    if dates1:
        sel.append("SELECT '1y', reference_date, source_sha256, COUNT(*) FROM observations_1y "
                   f"WHERE reference_date IN ({inlist(dates1)}) GROUP BY reference_date, source_sha256")
    out("  " + q.format(" UNION ALL ".join(sel)))
    out("  " + q.format("SELECT substr(sha256,1,12) AS sha, status, is_current, ingested_at FROM source_files "
                        f"WHERE reference_date IN ({inlist(sorted(set(dates5 + dates1)))}) "
                        "ORDER BY reference_date, is_current DESC"))
    out("  " + q.format("SELECT dataset_key, reference_date, kind FROM dataset_gaps ORDER BY dataset_key, reference_date"))
    out("  期待する値:")
    for (g, d), (sha, n) in sorted(expect.items()):
        out(f"    {g} {d}: {sha[:12]} で {n:,} 行")

    out("\n■ 3. 季節性の再計算とカタログ（前後で n_years の合計を比べる）")
    nq = q.format("SELECT 'pop' AS t, SUM(n_years) AS s, MIN(n_years) AS mn, MAX(n_years) AS mx FROM agg_seasonality_raw "
                  "UNION ALL SELECT 'hh', SUM(n_years), MIN(n_years), MAX(n_years) FROM agg_seasonality_hh_raw "
                  "UNION ALL SELECT 'age', SUM(n_years), MIN(n_years), MAX(n_years) FROM agg_seasonality_age_raw")
    out(f"  {nq}")
    out("  npx wrangler d1 time-travel info tokyo-population")
    for f in ("sql\\seasonality.sql", "sql\\seasonality_age.sql", "sql\\seasonality_households.sql",
              "scripts\\refresh_catalog.sql"):
        out(f"  npx wrangler d1 execute tokyo-population --remote --file {f}")
    out(f"  {nq}")
    out("  " + q.format("SELECT dataset_key, COUNT(*) AS periods, MIN(reference_date) AS first, MAX(reference_date) AS last "
                        "FROM dataset_periods WHERE dataset_key LIKE 'meguro%' GROUP BY dataset_key"))
    out("  " + q.format("SELECT dataset_key, SUM(has_residual) AS dates_with_residual, COUNT(*) AS dates "
                        "FROM dataset_sex_residual GROUP BY dataset_key"))
    out("  " + q.format("SELECT p.area_name, p.amplitude_pct AS pop_amp, h.amplitude_pct AS hh_amp, "
                        "p.trough_month AS pop_trough, h.trough_month AS hh_trough FROM v_nl_seasonality_strength p "
                        "JOIN v_nl_seasonality_hh_strength h ON h.key_code = p.key_code "
                        "ORDER BY p.amplitude_pct DESC LIMIT 5"))


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("step", choices=["check", "extract"])
    ap.add_argument("--month", required=True, help="対象の月（YYYY-MM）")
    ap.add_argument("--extra", action="append", default=[], help="追加で取り込む基準日（YYYY-MM-01）")
    ap.add_argument("--hold", action="append", default=None,
                    help=f"取り込まない基準日（既定 {DEFAULT_HOLD}。指定すると既定を置き換える）")
    args = ap.parse_args()
    return cmd_check(args) if args.step == "check" else cmd_extract(args)


if __name__ == "__main__":
    sys.exit(main())
